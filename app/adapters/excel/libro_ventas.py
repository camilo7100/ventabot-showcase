"""El Excel de ventas: una vista de la base, con columnas de Camilo (ADR 0015).

Se escribe ENTERO en cada sincronización, a un archivo temporal que luego
reemplaza al de verdad: un corte a mitad nunca deja un Excel roto. Si Camilo
lo tiene abierto, Windows no deja reemplazarlo; no pasa nada, se reintenta en
la siguiente vuelta y lo suyo sigue guardado en la base.

Las columnas del sistema quedan BLOQUEADAS (la hoja se protege sin clave:
se puede desproteger desde Revisar, pero no se edita por accidente). Las de
Camilo quedan libres y en verde.
"""

# ruff: noqa: E501  (SQL y textos del diccionario: partirlos los hace ilegibles)

from __future__ import annotations

import os
from datetime import datetime, timedelta, timezone
from pathlib import Path

from app.domain.ventas import (
    CAMILO,
    COLUMNAS_CLIENTAS,
    COLUMNAS_VENTAS,
    EDITABLES,
    Columna,
    valor_editable,
)

COLOMBIA = timezone(timedelta(hours=-5), "America/Bogota")
GRIS = "D9D9D9"
VERDE = "C6EFCE"
AZUL = "1F4E78"
FORMATOS = {"pesos": '"$"#,##0', "fecha": "yyyy-mm-dd hh:mm", "numero": "0.#"}


def _local(valor):
    """Excel no entiende zonas horarias: todo en hora de Colombia, sin zona."""
    if isinstance(valor, datetime):
        if valor.tzinfo is not None:
            valor = valor.astimezone(COLOMBIA)
        return valor.replace(tzinfo=None)
    return valor


class LibroDeVentas:
    def __init__(self, ruta: Path) -> None:
        self.ruta = Path(ruta)

    # -- leer lo que Camilo escribió ----------------------------------------

    def leer_editables(self) -> dict[str, dict[str, str]]:
        if not self.ruta.is_file():
            return {}
        from openpyxl import load_workbook

        libro = load_workbook(self.ruta, read_only=True, data_only=True)
        try:
            if "Ventas" not in libro.sheetnames:
                return {}
            hoja = libro["Ventas"]
            filas = hoja.iter_rows(values_only=True)
            titulos = next(filas, ())
            por_titulo = {c.titulo: c.clave for c in COLUMNAS_VENTAS}
            indice = {por_titulo[t]: i for i, t in enumerate(titulos) if t in por_titulo}
            if "id_venta" not in indice:
                return {}
            salida: dict[str, dict[str, str]] = {}
            for fila in filas:
                id_venta = (
                    fila[indice["id_venta"]] if len(fila) > indice["id_venta"] else None
                )
                if not id_venta:
                    continue
                salida[str(id_venta)] = {
                    clave: valor_editable(fila[indice[clave]])
                    for clave in EDITABLES
                    if clave in indice and len(fila) > indice[clave]
                }
            return salida
        finally:
            libro.close()

    # -- escribir la vista entera -------------------------------------------

    def escribir(
        self, ventas: list[dict], clientas: list[dict], registro: list[dict]
    ) -> None:
        from openpyxl import Workbook

        libro = Workbook()
        self._hoja(libro.active, "Ventas", COLUMNAS_VENTAS, ventas)
        self._hoja(libro.create_sheet(), "Clientas", COLUMNAS_CLIENTAS, clientas)
        self._diccionario(libro.create_sheet("Diccionario"))
        self._registro(libro.create_sheet("Registro"), registro)

        self.ruta.parent.mkdir(parents=True, exist_ok=True)
        temporal = self.ruta.with_name(f"~{self.ruta.stem}.tmp.xlsx")
        libro.save(temporal)
        try:
            os.replace(temporal, self.ruta)
        except PermissionError:
            temporal.unlink(missing_ok=True)
            raise

    def _hoja(self, hoja, titulo: str, columnas: tuple[Columna, ...], filas: list[dict]):
        from openpyxl.styles import Alignment, Font, PatternFill, Protection
        from openpyxl.utils import get_column_letter

        hoja.title = titulo
        for j, col in enumerate(columnas, 1):
            celda = hoja.cell(row=1, column=j, value=col.titulo)
            celda.font = Font(bold=True, color="000000")
            celda.fill = PatternFill(
                "solid", fgColor=VERDE if col.dueno == CAMILO else GRIS
            )
            celda.alignment = Alignment(wrap_text=True, vertical="center")
            ancho = 34 if col.clave in ("id_venta", "id_clienta", "notas_camilo") else 16
            hoja.column_dimensions[get_column_letter(j)].width = ancho
        for i, fila in enumerate(filas, 2):
            for j, col in enumerate(columnas, 1):
                celda = hoja.cell(row=i, column=j, value=_local(fila.get(col.clave)))
                if col.tipo in FORMATOS:
                    celda.number_format = FORMATOS[col.tipo]
                if col.dueno == CAMILO:
                    celda.protection = Protection(locked=False)
                    celda.fill = PatternFill("solid", fgColor="EAF7EC")
        hoja.freeze_panes = "C2"
        hoja.auto_filter.ref = hoja.dimensions
        hoja.protection.sheet = True
        # Proteger no es impedir mirar: filtrar, ordenar y ajustar columnas sí.
        hoja.protection.autoFilter = False
        hoja.protection.sort = False
        hoja.protection.formatColumns = False

    def _diccionario(self, hoja) -> None:
        from openpyxl.styles import Font

        hoja.append(["DICCIONARIO DE DATOS — Luna Joyas"])
        hoja["A1"].font = Font(bold=True, size=14, color=AZUL)
        for linea in (
            "La verdad está en la base de datos del sistema. Este archivo se regenera solo.",
            "Columnas GRISES: las escribe el sistema (lo que cambies ahí se pierde).",
            "Columnas VERDES: son tuyas. El sistema las guarda en la base y nunca las pisa.",
            "ID venta e ID clienta no cambian nunca: úsalos para cruzar con otros archivos.",
            "Contiene datos personales (Ley 1581 de 2012): no lo compartas ni lo subas sin cuidarlo.",
        ):
            hoja.append([linea])
        hoja.append([])
        hoja.append(["Hoja", "Columna", "Dueño", "Tipo", "Descripción"])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
        for nombre, columnas in (
            ("Ventas", COLUMNAS_VENTAS),
            ("Clientas", COLUMNAS_CLIENTAS),
        ):
            for c in columnas:
                hoja.append([nombre, c.titulo, c.dueno, c.tipo, c.descripcion])
        for letra, ancho in zip("ABCDE", (10, 24, 10, 10, 90), strict=True):
            hoja.column_dimensions[letra].width = ancho

    def _registro(self, hoja, registro: list[dict]) -> None:
        from openpyxl.styles import Font

        hoja.append(
            [
                "Fecha (UTC)",
                "Ventas",
                "Clientas",
                "Anotaciones tuyas guardadas",
                "Resultado",
            ]
        )
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for r in registro:
            hoja.append(
                [
                    r["created_at"],
                    r["ventas"],
                    r["clientas"],
                    r["anotaciones_nuevas"],
                    r["resultado"],
                ]
            )
        for letra, ancho in zip("ABCDE", (22, 8, 8, 26, 60), strict=True):
            hoja.column_dimensions[letra].width = ancho
