"""El Excel para Lía (cuestionario + lo del mes) y cómo entra al ERP (02/10).

    vb erp --cuestionario ARCHIVO.xlsx   crea el Excel, con lo que el ERP supone hoy
    vb erp --importar-lia ARCHIVO.xlsx   lo pasa al ERP (no duplica) y lo recalcula

El bot llena Ventas, Pagos y Produccion solo. Lo demás —lo que se compra, lo
que se paga, el precio de la plata, los equipos, lo que se vende por fuera de
WhatsApp, el conteo del inventario— lo anota Lía en este Excel, cada mes, y
este módulo lo copia a sus hojas. El cuestionario reemplaza los supuestos de
Costeo con lo que ella sabe; un «no sé» deja el supuesto como está.

Cada fila importada lleva una marca «LIA …» (en la nota, el motivo o la
referencia): importar el mismo archivo dos veces no duplica nada.
"""

from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path

from app.adapters.erp.montaje.comun import PIEZAS
from app.adapters.erp.montaje.costeo import SUPUESTOS, celda
from app.adapters.erp.montaje.gastos import GASTOS_DE_LIA
from app.adapters.erp.sheets import Hoja, _letra
from app.domain import erp
from app.domain.cuestionario import (
    NO_SE,
    PREGUNTAS,
    Respuestas,
    cambios_del_costeo,
    leer_pesos,
)

RESPUESTAS = "Respuestas_Lía"
MARCA = "LIA"
CANALES_DE_VENTA = ("En persona", "Feria", "Página web", "Instagram", "Otro")
TIPOS_DE_AJUSTE = "Conteo físico"

#: Las hojas del Excel: (nombre, encabezados, ancho de cada columna).
HOJAS = (
    (
        "Compras",
        ("Fecha", "Proveedor", "Material", "Cantidad", "Valor total ($)", "Nota"),
        (12, 22, 28, 11, 16, 40),
    ),
    ("Gastos", ("Fecha", "Concepto", "Valor ($)", "Nota"), (12, 44, 16, 44)),
    ("Plata", ("Mes (AAAA-MM)", "Precio por gramo ($)", "Nota"), (16, 20, 50)),
    (
        "Equipos",
        ("Equipo", "Año de compra", "Costo ($)", "Años que dura", "Nota"),
        (34, 14, 16, 14, 40),
    ),
    (
        "Ventas fuera del bot",
        ("Fecha", "Producto", "Cantidad", "Valor total ($)", "Canal", "Ciudad", "Nota"),
        (12, 22, 10, 16, 16, 18, 36),
    ),
    (
        "Inventario hoy",
        ("Material", "Cantidad contada", "Fecha del conteo", "Nota"),
        (30, 18, 16, 40),
    ),
)


# ══════════════════════════════════════════════════════════════════════════
#  1. EL EXCEL PARA LIA
# ══════════════════════════════════════════════════════════════════════════


def crear_libro(destino: Path | io.BytesIO, materiales: list[str], fecha: str) -> None:
    """El Excel que se le pasa a Lía: Léeme, Preguntas y las hojas del mes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    amarillo = PatternFill("solid", fgColor="FFF2CC")
    gris = PatternFill("solid", fgColor="EDEDED")
    titulo = Font(bold=True, size=14)
    negrita = Font(bold=True)
    envolver = Alignment(wrap_text=True, vertical="top")

    libro = Workbook()
    leeme = libro.active
    leeme.title = "Léeme"
    lineas = [
        ("Luna Joyas · lo que el ERP todavía no sabe", titulo),
        (f"Preparado el {fecha} por VentaBot (Camilo).", None),
        ("", None),
        (
            "1. Hoja «Preguntas»: contesta en la columna amarilla. Si no sabes, "
            f"escribe {NO_SE}. No dejes ninguna en blanco.",
            None,
        ),
        (
            "   Escribe los valores como te salga: $1.200.000 · 1,2 millones · "
            "600 mil · 8 %.",
            None,
        ),
        (
            "2. Las otras hojas son de CADA MES (lo que el bot no ve). Una fila por "
            "cosa, con su fecha. Si un mes no hubo, déjala vacía.",
            None,
        ),
        (
            "   · Compras: plata, piedras, empaques, insumos. El valor es lo que "
            "pagaste en total.",
            None,
        ),
        (
            "   · Gastos: arriendo, servicios, pagos a la familia, taxis… Si anotas "
            "algo del TALLER un mes, anota TODO lo del taller de ese mes (ese mes el "
            "ERP deja de usar el estimado).",
            None,
        ),
        ("   · Plata: el precio por gramo al que la compraste cada mes.", None),
        ("   · Equipos: una vez, la lista de máquinas con lo que costaron.", None),
        ("   · Ventas fuera del bot: lo que vendiste por fuera de WhatsApp.", None),
        ("   · Inventario hoy: cuando cuentes lo que hay guardado.", None),
        (
            "3. Devuélvele el archivo a Camilo. Se puede importar varias veces: lo "
            "que ya entró no se duplica (si corriges una fila ya importada, avísale).",
            None,
        ),
    ]
    for i, (texto, fuente) in enumerate(lineas, 1):
        c = leeme.cell(row=i, column=1, value=texto)
        if fuente:
            c.font = fuente
    leeme.column_dimensions["A"].width = 110

    preguntas = libro.create_sheet("Preguntas")
    cabeza = (
        "#",
        "Tema",
        "Pregunta",
        "Para qué sirve",
        "Cómo se contesta",
        "TU RESPUESTA",
    )
    for j, (texto, ancho) in enumerate(
        zip(cabeza, (5, 18, 60, 42, 20, 34), strict=True), 1
    ):
        c = preguntas.cell(row=1, column=j, value=texto)
        c.font, c.fill = negrita, gris
        preguntas.column_dimensions[_letra(j - 1)].width = ancho
    for i, p in enumerate(PREGUNTAS, 2):
        for j, valor in enumerate((p.id, p.tema, p.texto, p.para_que, p.unidad, ""), 1):
            c = preguntas.cell(row=i, column=j, value=valor)
            c.alignment = envolver
        preguntas.cell(row=i, column=6).fill = amarillo
    preguntas.freeze_panes = "A2"

    listas = libro.create_sheet("Listas")
    columnas_de_listas = {
        "materiales": materiales,
        "conceptos": [c for c, _, _ in GASTOS_DE_LIA],
        "productos": [*PIEZAS, "Dije", "Manilla", "Otro"],
        "canales": list(CANALES_DE_VENTA),
    }
    rangos = {}
    for j, (nombre, valores) in enumerate(columnas_de_listas.items(), 1):
        listas.cell(row=1, column=j, value=nombre)
        for i, v in enumerate(valores, 2):
            listas.cell(row=i, column=j, value=v)
        letra = _letra(j - 1)
        rangos[nombre] = f"Listas!${letra}$2:${letra}${len(valores) + 1}"
    listas.sheet_state = "hidden"

    con_lista = {
        ("Compras", "Material"): "materiales",
        ("Gastos", "Concepto"): "conceptos",
        ("Ventas fuera del bot", "Producto"): "productos",
        ("Ventas fuera del bot", "Canal"): "canales",
        ("Inventario hoy", "Material"): "materiales",
    }
    for nombre, encabezados, anchos in HOJAS:
        h = libro.create_sheet(nombre)
        for j, (texto, ancho) in enumerate(zip(encabezados, anchos, strict=True), 1):
            c = h.cell(row=1, column=j, value=texto)
            c.font, c.fill = negrita, gris
            h.column_dimensions[_letra(j - 1)].width = ancho
            for i in range(2, 202):
                h.cell(row=i, column=j).fill = amarillo
            if (lista := con_lista.get((nombre, texto))) is not None:
                v = DataValidation(type="list", formula1=f"={rangos[lista]}")
                v.error = "Escoge de la lista (o escribe a Camilo si falta algo)"
                h.add_data_validation(v)
                v.add(f"{_letra(j - 1)}2:{_letra(j - 1)}201")
        h.freeze_panes = "A2"
    libro.save(destino)


# ══════════════════════════════════════════════════════════════════════════
#  2. LEERLO
# ══════════════════════════════════════════════════════════════════════════


@dataclass
class LibroDeLia:
    respuestas: Respuestas = field(default_factory=Respuestas)
    #: hoja -> filas (dict encabezado -> valor), sin las vacías.
    filas: dict[str, list[dict]] = field(default_factory=dict)


def leer_libro(datos: bytes) -> LibroDeLia:
    import openpyxl

    libro = openpyxl.load_workbook(io.BytesIO(datos), data_only=True)
    salida = LibroDeLia()
    if "Preguntas" in libro.sheetnames:
        for fila in libro["Preguntas"].iter_rows(min_row=2, values_only=True):
            if fila and fila[0]:
                salida.respuestas.escritas[str(fila[0]).strip()] = (
                    "" if len(fila) < 6 or fila[5] is None else str(fila[5]).strip()
                )
    for nombre, encabezados, _ in HOJAS:
        if nombre not in libro.sheetnames:
            continue
        filas = []
        for i, fila in enumerate(libro[nombre].iter_rows(min_row=2, values_only=True), 2):
            if not fila or all(v in (None, "") for v in fila):
                continue
            d = {
                e: (fila[j] if j < len(fila) else None) for j, e in enumerate(encabezados)
            }
            d["_fila"] = i
            filas.append(d)
        salida.filas[nombre] = filas
    return salida


def _fecha(valor) -> str | None:
    """La fecha como AAAA-MM-DD, venga como fecha de Excel o como texto."""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return (
                datetime.strptime(texto, formato).replace(tzinfo=UTC).date().isoformat()
            )
        except ValueError:
            continue
    return None


def _mes(valor) -> int | None:
    """«2026-10», «10/2026», una fecha o 202610 -> 202610."""
    if isinstance(valor, datetime | date):
        return valor.year * 100 + valor.month
    texto = str(valor or "").strip().replace("/", "-")
    partes = [p for p in texto.split("-") if p.isdigit()]
    if len(partes) == 2:
        a, m = (int(p) for p in partes)
        if a < 100:
            a, m = m, a
        return a * 100 + m if 1 <= m <= 12 and 2000 <= a <= 2099 else None
    if texto.isdigit() and len(texto) == 6:
        return int(texto)
    return None


def _marca(hoja: str, *valores) -> str:
    huella = hashlib.sha256("|".join(str(v) for v in valores).encode()).hexdigest()[:8]
    return f"{MARCA} {hoja} {huella}"


# ══════════════════════════════════════════════════════════════════════════
#  3. PASARLO AL ERP
# ══════════════════════════════════════════════════════════════════════════


@dataclass
class Importado:
    hechos: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


def _tabla(hoja: Hoja, nombre: str, columnas: str = "A3:AH3000"):
    (filas,) = hoja.leer([f"{nombre}!{columnas}"])
    encabezados = [str(x).strip() for x in (filas[0] if filas else [])]
    return encabezados, filas[1:]


def _libres(filas: list[list], desde: int = 4):
    """Filas sin fecha/nombre en la columna A (las que ya traen fórmulas, primero)."""
    for i, f in enumerate(filas):
        if not (f and str(f[0]).strip()):
            yield desde + i
    n = desde + len(filas)
    while True:
        yield n
        n += 1


def pasar_al_erp(hoja: Hoja, libro: LibroDeLia, fecha: str) -> Importado:
    r = Importado()
    e: list[tuple[str, list[list]]] = []
    _cuestionario(hoja, libro, fecha, e, r)
    _compras(hoja, libro, e, r)
    _gastos(hoja, libro, e, r)
    _plata(hoja, libro, fecha, e, r)
    _equipos(hoja, libro, e, r)
    _ventas(hoja, libro, e, r)
    _inventario(hoja, libro, e, r)
    hoja.escribir(e)
    return r


def _cuestionario(hoja, libro, fecha, e, r) -> None:
    respuestas = libro.respuestas
    if not respuestas.escritas:
        return
    hoja.asegurar_hoja(RESPUESTAS)
    filas = [["#", "Tema", "Pregunta", "Respuesta de Lía", "Anotada"]]
    for p in PREGUNTAS:
        filas.append([p.id, p.tema, p.texto, respuestas.escritas.get(p.id, ""), fecha])
    e.append((f"{RESPUESTAS}!A1", filas))
    faltan = respuestas.sin_contestar()
    if faltan:
        r.avisos.append(
            f"{len(faltan)} preguntas en blanco: " + ", ".join(p.id for p in faltan)
        )
    nombres = {s.clave: s.nombre for s in SUPUESTOS}
    (materiales, config) = hoja.leer(["Cat_Materiales!A1:H60", "Config!A1:F20"])
    for cambio in cambios_del_costeo(respuestas, fecha):
        if cambio.destino in nombres:
            fila = celda(f"S:{cambio.destino}").split("$")[-1]
            e.append((f"Costeo!B{fila}", [[cambio.valor]]))
            e.append((f"Costeo!D{fila}", [["REAL", cambio.fuente]]))
            r.hechos.append(f"Costeo: {nombres[cambio.destino]} = {cambio.valor:,.0f}")
        elif cambio.destino.startswith("material:"):
            material = cambio.destino.split(":", 1)[1]
            n = _fila_con(materiales, material)
            if n is None:
                r.avisos.append(f"el material «{material}» no está en Cat_Materiales")
                continue
            e.append((f"Cat_Materiales!C{n}", [[cambio.valor]]))
            e.append((f"Cat_Materiales!F{n}", [["REAL", cambio.fuente, fecha]]))
            r.hechos.append(f"{material}: ${cambio.valor:,.0f}")
        elif cambio.destino == "mano_de_obra":
            n = _fila_con(config, "Mano de obra por pieza ($)")
            if n is not None:
                e.append((f"Config!B{n}", [[cambio.valor]]))
                e.append((f"Config!D{n}", [["REAL", cambio.fuente, fecha]]))
                r.hechos.append(f"mano de obra por pieza: ${cambio.valor:,.0f}")


def _fila_con(filas: list[list], nombre: str) -> int | None:
    for i, f in enumerate(filas, 1):
        if f and str(f[0]).strip() == nombre:
            return i
    return None


def _ya(filas: list[list], columna: int) -> set[str]:
    return {
        str(f[columna]).split(" · ")[0].strip()
        for f in filas
        if len(f) > columna and str(f[columna]).startswith(MARCA)
    }


def _compras(hoja, libro, e, r) -> None:
    filas_lia = libro.filas.get("Compras", [])
    if not filas_lia:
        return
    _, filas = _tabla(hoja, "Compras", "A3:I3000")
    hechas, libres, nuevas = _ya(filas, 7), _libres(filas), 0
    for f in filas_lia:
        dia, material = _fecha(f["Fecha"]), str(f["Material"] or "").strip()
        cantidad, total = leer_pesos(f["Cantidad"]), leer_pesos(f["Valor total ($)"])
        if not (dia and material and cantidad and total):
            r.avisos.append(
                f"Compras fila {f['_fila']}: falta fecha, material, "
                "cantidad o valor (no se importó)"
            )
            continue
        marca = _marca("compra", dia, material, cantidad, total)
        if marca in hechas:
            continue
        n = next(libres)
        e += [
            (f"Compras!A{n}", [[dia, f["Proveedor"] or "", material, cantidad]]),
            (
                f"Compras!E{n}",
                [
                    [
                        f'=IF($C{n}="";"";IFERROR(INDEX(Cat_Materiales!$B$4:$B$56;'
                        f'MATCH($C{n};Cat_Materiales!$A$4:$A$56;0));""))',
                        round(total / cantidad, 2),
                        f'=IF($C{n}="";"";$D{n}*$F{n})',
                        f"{marca} · {f['Nota'] or ''}".strip(" ·"),
                        f'=IF($A{n}="";"";YEAR($A{n})*100+MONTH($A{n}))',
                    ]
                ],
            ),
        ]
        hechas.add(marca)
        nuevas += 1
    if nuevas:
        r.hechos.append(f"Compras: +{nuevas}")


def _gastos(hoja, libro, e, r) -> None:
    filas_lia = libro.filas.get("Gastos", [])
    if not filas_lia:
        return
    categoria = {c: cat for c, cat, _ in GASTOS_DE_LIA}
    _, filas = _tabla(hoja, "Gastos", "A3:G3000")
    hechas, libres, nuevas = _ya(filas, 5), _libres(filas), 0
    for f in filas_lia:
        dia, concepto = _fecha(f["Fecha"]), str(f["Concepto"] or "").strip()
        valor = leer_pesos(f["Valor ($)"])
        if not (dia and concepto and valor is not None):
            r.avisos.append(f"Gastos fila {f['_fila']}: falta fecha, concepto o valor")
            continue
        marca = _marca("gasto", dia, concepto, valor)
        if marca in hechas:
            continue
        n = next(libres)
        e.append(
            (
                f"Gastos!A{n}",
                [
                    [
                        dia,
                        concepto,
                        categoria.get(concepto, "Otro"),
                        valor,
                        "REAL",
                        f"{marca} · {f['Nota'] or ''}".strip(" ·"),
                        f'=IF($A{n}="";"";YEAR($A{n})*100+MONTH($A{n}))',
                    ]
                ],
            )
        )
        hechas.add(marca)
        nuevas += 1
    if nuevas:
        r.hechos.append(f"Gastos: +{nuevas} (el estimado de esos meses se ajusta solo)")


def _plata(hoja, libro, fecha, e, r) -> None:
    filas_lia = libro.filas.get("Plata", [])
    if not filas_lia:
        return
    (filas,) = hoja.leer(["Precio_Plata!A4:D200"])
    donde = {str(f[0]).strip(): 4 + i for i, f in enumerate(filas) if f}
    siguiente = 4 + len(filas)
    for f in filas_lia:
        mes, precio = _mes(f["Mes (AAAA-MM)"]), leer_pesos(f["Precio por gramo ($)"])
        if not (mes and precio):
            r.avisos.append(f"Plata fila {f['_fila']}: falta el mes o el precio")
            continue
        n = donde.get(str(mes))
        if n is None:
            n, siguiente = siguiente, siguiente + 1
        fuente = f"Lía, {fecha}" + (f" · {f['Nota']}" if f["Nota"] else "")
        e.append((f"Precio_Plata!A{n}", [[mes, precio, fuente, fecha]]))
        r.hechos.append(f"plata {mes}: ${precio:,.0f}/g")


def _equipos(hoja, libro, e, r) -> None:
    filas_lia = libro.filas.get("Equipos", [])
    if not filas_lia:
        return
    _, filas = _tabla(hoja, "Activos", "A3:I200")
    hechas, libres, nuevas = _ya(filas, 8), _libres(filas), 0
    for f in filas_lia:
        nombre, costo = str(f["Equipo"] or "").strip(), leer_pesos(f["Costo ($)"])
        anio = leer_pesos(f["Año de compra"])
        anos = leer_pesos(f["Años que dura"]) or 10
        if not (nombre and costo):
            r.avisos.append(f"Equipos fila {f['_fila']}: falta el equipo o el costo")
            continue
        marca = _marca("equipo", nombre, costo, anio)
        if marca in hechas:
            continue
        n = next(libres)
        compra = f"{int(anio)}-01-01" if anio and 1990 <= anio <= 2099 else ""
        e.append(
            (
                f"Activos!A{n}",
                [
                    [
                        nombre,
                        compra,
                        costo,
                        int(anos * 12),
                        f'=IF(OR($A{n}="";$D{n}=0;$D{n}="");"";$C{n}/$D{n})',
                        "SI",
                        "REAL",
                        "Excel de datos (vb erp --importar-lia)",
                        f"{marca} · {f['Nota'] or ''}".strip(" ·"),
                    ]
                ],
            )
        )
        hechas.add(marca)
        nuevas += 1
    if nuevas:
        r.hechos.append(f"Equipos: +{nuevas}")
        # El bloque estimado de $30 M deja de contar: ya están los de verdad.
        n = _fila_con([[""]] * 2 + filas, "Taller y equipos (consolidado)")
        if n is not None:
            e.append((f"Activos!F{n + 1}", [["NO"]]))
            r.hechos.append("el estimado «Taller y equipos (consolidado)» queda en NO")


def _ventas(hoja, libro, e, r) -> None:
    filas_lia = libro.filas.get("Ventas fuera del bot", [])
    if not filas_lia:
        return
    encabezados, filas = _tabla(hoja, "Ventas", "A3:AH3000")
    if erp.REFERENCIA not in encabezados:
        r.avisos.append("Ventas no tiene la columna Pedido_Ref: no se importaron ventas")
        return
    c_ref = encabezados.index(erp.REFERENCIA)
    hechas = {str(f[c_ref]).strip() for f in filas if len(f) > c_ref}
    ids = [str(f[0]) for f in filas if f]
    libres, nuevas = _libres(filas), 0
    for f in filas_lia:
        dia, producto = _fecha(f["Fecha"]), str(f["Producto"] or "").strip()
        total = leer_pesos(f["Valor total ($)"])
        cantidad = int(leer_pesos(f["Cantidad"]) or 1)
        if not (dia and producto and total):
            r.avisos.append(f"Ventas fila {f['_fila']}: falta fecha, producto o valor")
            continue
        marca = _marca("venta", dia, producto, cantidad, total)
        if marca in hechas:
            continue
        n = next(libres)
        venta_id = erp.siguiente_venta_id(ids)
        ids.append(venta_id)
        e.append(
            (
                f"Ventas!A{n}",
                [
                    [
                        venta_id,
                        dia,
                        producto,
                        "",
                        f["Canal"] or "Otro",
                        f["Ciudad"] or "",
                        "",
                        "",
                        str(f["Nota"] or ""),
                        cantidad,
                        round(total / cantidad),
                        total,
                        "",
                        "Entregado",
                    ]
                ],
            )
        )
        e.append((f"Ventas!{_letra(c_ref)}{n}", [[marca]]))
        hechas.add(marca)
        nuevas += 1
    if nuevas:
        r.hechos.append(f"Ventas fuera del bot: +{nuevas}")


def _inventario(hoja, libro, e, r) -> None:
    filas_lia = libro.filas.get("Inventario hoy", [])
    if not filas_lia:
        return
    (inventario,) = hoja.leer(["Inventario!A4:F80"])
    stock = {
        str(f[0]).strip(): leer_pesos(f[5]) or 0
        for f in inventario
        if len(f) > 5 and str(f[0]).strip()
    }
    _, filas = _tabla(hoja, "Ajustes", "A3:G3000")
    hechas, libres, nuevas = _ya(filas, 5), _libres(filas), 0
    for f in filas_lia:
        material = str(f["Material"] or "").strip()
        contado = leer_pesos(f["Cantidad contada"])
        dia = _fecha(f["Fecha del conteo"]) or datetime.now(UTC).date().isoformat()
        if not material or contado is None:
            r.avisos.append(f"Inventario fila {f['_fila']}: falta material o cantidad")
            continue
        if material not in stock:
            r.avisos.append(f"Inventario: «{material}» no está en el ERP")
            continue
        marca = _marca("conteo", dia, material, contado)
        if marca in hechas:
            continue
        diferencia = contado - stock[material]
        n = next(libres)
        e.append(
            (
                f"Ajustes!A{n}",
                [
                    [
                        dia,
                        TIPOS_DE_AJUSTE,
                        material,
                        round(diferencia, 3),
                        "",
                        f"{marca} · contó {contado:g}; el ERP tenía {stock[material]:g}",
                        f'=IF($A{n}="";"";YEAR($A{n})*100+MONTH($A{n}))',
                    ]
                ],
            )
        )
        hechas.add(marca)
        nuevas += 1
    if nuevas:
        r.hechos.append(f"Inventario: {nuevas} conteos (como ajustes)")
